from playwright.sync_api import sync_playwright
from openpyxl import load_workbook
from datetime import datetime
import time

# ==========================================
# CONFIGURAÇÕES
# ==========================================

ARQUIVO_EXCEL = "Levantamento-semanal-de-multas.xlsx"
ABA = "AUTOPRF"
URL_PRF = "https://pesquisa-auto.prf.gov.br/#/pesquisa/consultar-debitos"

COL_PLACA = 2     # coluna B
COL_RENAVAM = 4   # coluna D
TEXTO_SEM_MULTA = "Nenhum registro encontrado"


# ==========================================
# FUNÇÕES AUXILIARES
# ==========================================

def salvar(workbook):
    """Tenta salvar; se o Excel estiver aberto, espera e tenta de novo."""
    while True:
        try:
            workbook.save(ARQUIVO_EXCEL)
            return
        except PermissionError:
            print("!! Não consegui salvar. Feche o Excel (o arquivo está em uso).")
            print("   Tentando novamente em 10 segundos...")
            time.sleep(10)


def limpar_placa(valor):
    return str(valor).strip().upper().replace("-", "").replace(" ", "")


def limpar_renavam(valor):
    texto = str(valor).strip()
    if texto.endswith(".0"):
        texto = texto[:-2]
    # se o Excel guardou como número, recupera os zeros à esquerda
    if texto.isdigit() and len(texto) < 11:
        texto = texto.zfill(11)
    return texto


# ==========================================
# ABRIR PLANILHA
# ==========================================

workbook = load_workbook(ARQUIVO_EXCEL)
planilha = workbook[ABA]

# ==========================================
# COLUNA DA DATA DE HOJE
# ==========================================

cabecalho_data = f"DATA: {datetime.now().strftime('%d/%m/%Y')}"

coluna_data = None
ultima_coluna_com_cabecalho = 0

for col in range(1, planilha.max_column + 1):
    valor = planilha.cell(1, col).value
    if valor is not None and str(valor).strip() != "":
        ultima_coluna_com_cabecalho = col
        if str(valor).strip() == cabecalho_data:
            coluna_data = col

if coluna_data is None:
    coluna_data = ultima_coluna_com_cabecalho + 1
    planilha.cell(1, coluna_data).value = cabecalho_data

print("-----------------------------------")
print(cabecalho_data, f"(coluna {coluna_data})")
print("-----------------------------------")

salvar(workbook)

# ==========================================
# ÚLTIMA LINHA REAL (com placa)
# ==========================================

ultima_linha = 1
for l in range(2, planilha.max_row + 1):
    if planilha.cell(l, COL_PLACA).value:
        ultima_linha = l

total = ultima_linha - 1
print(f"Total de veículos: {total}")

# ==========================================
# NAVEGADOR
# ==========================================

with sync_playwright() as p:

    browser = p.chromium.launch(headless=False)
    page = browser.new_page()

    for linha in range(2, ultima_linha + 1):

        placa_raw = planilha.cell(linha, COL_PLACA).value
        renavam_raw = planilha.cell(linha, COL_RENAVAM).value

        if not placa_raw or not renavam_raw:
            print(f"Linha {linha}: sem placa/RENAVAM. Pulando...")
            continue

        # pula o que já foi consultado hoje (permite retomar)
        ja_feito = planilha.cell(linha, coluna_data).value
        if ja_feito in ("OK", "MULTA"):
            print(f"Linha {linha}: já consultada hoje ({ja_feito}). Pulando...")
            continue

        placa = limpar_placa(placa_raw)
        renavam = limpar_renavam(renavam_raw)

        print()
        print("===================================")
        print(f"CONSULTA {linha - 1} DE {total}")
        print(f"Placa: {placa} | RENAVAM: {renavam}")
        print("===================================")

        try:
            page.goto(URL_PRF, wait_until="domcontentloaded")

            campos = page.locator('input:not([type="checkbox"])')
            campos.first.wait_for(state="visible", timeout=20000)

            campos.nth(0).fill(placa)
            campos.nth(1).fill(renavam)

            page.get_by_role("button", name="Pesquisar").click()

            # espera aparecer OU a mensagem de "sem registro" OU linhas na tabela
            resultado = page.get_by_text(TEXTO_SEM_MULTA).or_(
                page.locator("table tbody tr")
            ).first
            resultado.wait_for(state="visible", timeout=30000)

            if page.get_by_text(TEXTO_SEM_MULTA).count() > 0:
                print("SEM MULTA -> OK")
                planilha.cell(linha, coluna_data).value = "OK"
            else:
                print("REGISTRO ENCONTRADO -> MULTA")
                planilha.cell(linha, coluna_data).value = "MULTA"

        except Exception as erro:
            print("ERRO NA CONSULTA:", str(erro).splitlines()[0])
            planilha.cell(linha, coluna_data).value = "ERRO"

        salvar(workbook)
        print("Salvo na planilha.")

    print()
    print("===================================")
    print("TODAS AS CONSULTAS FORAM FINALIZADAS")
    print("===================================")

    input("\nPressione ENTER para fechar o navegador...")
    browser.close()